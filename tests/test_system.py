import os
import sys
import unittest
from datetime import datetime
import openpyxl

# Add backend to path
BACKEND_DIR = os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))), "backend")
sys.path.insert(0, BACKEND_DIR)

from fastapi.testclient import TestClient
from main import app
from services.auth_service import auth_service, STORE_FILE
from services.excel_service import EXCEL_FILE_PATH


class TestVisitorEntrySystem(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        # Reset auth store for clean testing
        if os.path.exists(STORE_FILE):
            try:
                os.remove(STORE_FILE)
            except Exception:
                pass
        auth_service.__init__()
        cls.client = TestClient(app)

    def test_01_health_check(self):
        response = self.client.get("/api/health")
        self.assertEqual(response.status_code, 200)
        data = response.json()
        self.assertEqual(data["status"], "healthy")

    def test_02_authorize_invalid_pin(self):
        payload = {
            "visitor_names": ["Rahul Sharma", "Priya Patil"],
            "pin": "9999"  # Incorrect PIN
        }
        response = self.client.post("/api/authorize-entry", json=payload)
        self.assertEqual(response.status_code, 200)
        data = response.json()
        self.assertFalse(data["success"])
        self.assertEqual(data["authorization_status"], "NOT AUTHORIZED")

    def test_03_authorize_empty_names(self):
        payload = {
            "visitor_names": ["   "],
            "pin": "1234"
        }
        response = self.client.post("/api/authorize-entry", json=payload)
        self.assertEqual(response.status_code, 422)

    def test_04_authorize_success(self):
        payload = {
            "visitor_names": ["Rahul Sharma", "Priya Patil"],
            "pin": "1234"
        }
        response = self.client.post("/api/authorize-entry", json=payload)
        self.assertEqual(response.status_code, 200)
        data = response.json()
        self.assertTrue(data["success"])
        self.assertEqual(data["authorization_status"], "AUTHORIZED")
        self.assertEqual(data["visitor_names"], ["Rahul Sharma", "Priya Patil"])

    def test_05_gate_verify_success_and_excel(self):
        payload = {
            "visitor_names": ["Rahul Sharma", "Priya Patil"],
            "pin": "1234"
        }
        response = self.client.post("/api/verify-entry", json=payload)
        self.assertEqual(response.status_code, 200)
        data = response.json()
        self.assertTrue(data["success"])
        self.assertEqual(data["entry_status"], "ENTRY AUTHORIZED")

        # Verify Excel file exists and has the recorded row
        self.assertTrue(os.path.exists(EXCEL_FILE_PATH))
        wb = openpyxl.load_workbook(EXCEL_FILE_PATH)
        ws = wb.active
        self.assertGreaterEqual(ws.max_row, 2)
        row_vals = [ws.cell(row=ws.max_row, column=c).value for c in range(1, 6)]
        self.assertIn("Rahul Sharma", row_vals[0])
        self.assertIn("Priya Patil", row_vals[0])
        self.assertEqual(row_vals[3], "AUTHORIZED")
        self.assertEqual(row_vals[4], "ENTRY AUTHORIZED")

    def test_06_duplicate_entry_protection(self):
        payload = {
            "visitor_names": ["Rahul Sharma", "Priya Patil"],
            "pin": "1234"
        }
        # Attempt to verify the same entry again
        response = self.client.post("/api/verify-entry", json=payload)
        self.assertEqual(response.status_code, 200)
        data = response.json()
        self.assertFalse(data["success"])
        self.assertEqual(data["entry_status"], "ENTRY ALREADY COMPLETED")

    def test_07_gate_verify_unauthorized_visitor(self):
        payload = {
            "visitor_names": ["Unknown Stranger"],
            "pin": "1234"
        }
        response = self.client.post("/api/verify-entry", json=payload)
        self.assertEqual(response.status_code, 200)
        data = response.json()
        self.assertFalse(data["success"])
        self.assertEqual(data["entry_status"], "ENTRY NOT AUTHORIZED")

    def test_08_admin_change_pin(self):
        payload = {
            "current_pin": "1234",
            "new_pin": "5678",
            "confirm_new_pin": "5678",
            "admin_secret": "admin123"
        }
        response = self.client.post("/api/admin/change-pin", json=payload)
        self.assertEqual(response.status_code, 200)
        data = response.json()
        self.assertTrue(data["success"])

        # Check old PIN is now rejected
        old_attempt = self.client.post(
            "/api/authorize-entry",
            json={"visitor_names": ["Amit Kumar"], "pin": "1234"}
        )
        self.assertEqual(old_attempt.json()["authorization_status"], "NOT AUTHORIZED")

        # Check new PIN is now accepted
        new_attempt = self.client.post(
            "/api/authorize-entry",
            json={"visitor_names": ["Amit Kumar"], "pin": "5678"}
        )
        self.assertEqual(new_attempt.json()["authorization_status"], "AUTHORIZED")


if __name__ == "__main__":
    unittest.main()
