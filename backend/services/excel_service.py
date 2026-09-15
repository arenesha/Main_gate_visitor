import os
from typing import List, Union, Dict, Any
from datetime import datetime
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

BASE_DIR = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
DATA_DIR = os.path.join(BASE_DIR, "data")
EXCEL_FILE_PATH = os.path.join(DATA_DIR, "visitor_entries.xlsx")

HEADERS = [
    "Visitor Name(s)",
    "Entry Date",
    "Entry Time",
    "Authorization Status",
    "Entry Status"
]

def ensure_data_directory() -> None:
    os.makedirs(DATA_DIR, exist_ok=True)

def append_gate_log_entry(
    visitor_names: Union[List[str], str],
    entry_date: str,
    entry_time: str,
    authorization_status: str,
    entry_status: str,
    file_path: str = EXCEL_FILE_PATH
) -> str:
    """
    Appends a gate check record (ENTRY AUTHORIZED or ENTRY NOT AUTHORIZED) to the Excel log.
    """
    ensure_data_directory()

    names_str = ", ".join(visitor_names) if isinstance(visitor_names, list) else str(visitor_names)

    if not os.path.exists(file_path):
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Visitor Entries"
        ws.append(HEADERS)
    else:
        try:
            wb = openpyxl.load_workbook(file_path)
            ws = wb.active
        except Exception:
            wb = openpyxl.Workbook()
            ws = wb.active
            ws.title = "Visitor Entries"
            ws.append(HEADERS)

    ws.append([names_str, entry_date, entry_time, authorization_status, entry_status])

    # Formatting styling
    header_fill = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
    header_font = Font(name="Inter", size=11, bold=True, color="FFFFFF")
    center_align = Alignment(horizontal="center", vertical="center")
    thin_border = Border(
        left=Side(style="thin", color="CBD5E1"),
        right=Side(style="thin", color="CBD5E1"),
        top=Side(style="thin", color="CBD5E1"),
        bottom=Side(style="thin", color="CBD5E1")
    )

    for col_idx in range(1, len(HEADERS) + 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = center_align
        cell.border = thin_border

    ws.row_dimensions[1].height = 26

    # Data row styling
    new_row_idx = ws.max_row
    ws.row_dimensions[new_row_idx].height = 22
    regular_font = Font(name="Inter", size=10)

    for col_idx in range(1, len(HEADERS) + 1):
        cell = ws.cell(row=new_row_idx, column=col_idx)
        cell.font = regular_font
        cell.border = thin_border
        if col_idx in (2, 3, 4, 5):
            cell.alignment = center_align
        else:
            cell.alignment = Alignment(horizontal="left", vertical="center")

    # Column width auto-fit
    for col in ws.columns:
        max_len = max(len(str(c.value or "")) for c in col)
        col_letter = get_column_letter(col[0].column)
        ws.column_dimensions[col_letter].width = max(max_len + 4, 18)

    wb.save(file_path)
    return file_path

def read_excel_dashboard_data(file_path: str = EXCEL_FILE_PATH) -> Dict[str, Any]:
    """
    Reads the Excel file and returns statistics and recent log entries.
    """
    today_str = datetime.now().strftime("%Y-%m-%d")
    stats = {
        "checks_today": 0,
        "authorized_today": 0,
        "denied_today": 0,
        "recent_entries": []
    }

    if not os.path.exists(file_path):
        return stats

    try:
        wb = openpyxl.load_workbook(file_path, data_only=True)
        ws = wb.active
        rows = list(ws.iter_rows(values_only=True))

        if len(rows) <= 1:
            return stats

        data_rows = rows[1:]
        for r in data_rows:
            if not r or len(r) < 5:
                continue
            names, entry_date, entry_time, auth_status, entry_status = r[0], str(r[1]), str(r[2]), str(r[3]), str(r[4])

            if entry_date == today_str:
                stats["checks_today"] += 1
                if "AUTHORIZED" in entry_status and "NOT" not in entry_status:
                    stats["authorized_today"] += 1
                else:
                    stats["denied_today"] += 1

            stats["recent_entries"].append({
                "visitor_names": names,
                "entry_date": entry_date,
                "entry_time": entry_time,
                "authorization_status": auth_status,
                "entry_status": entry_status
            })

        # Reverse entries so newest appear first
        stats["recent_entries"].reverse()
    except Exception as e:
        pass

    return stats
